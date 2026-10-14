"""把同文件夹里的 Excel 网站数据导入 SQLite 数据库。

第一次使用前，在 VS Code 终端执行：python -m pip install openpyxl
运行：python SQL_practice/import_excel.py
"""

from pathlib import Path
import sqlite3

from openpyxl import load_workbook


# 使用脚本所在的文件夹定位文件，不依赖终端当前打开的目录。
folder = Path(__file__).resolve().parent
excel_path = folder / "tech_sites_practice.xlsx"
db_path = folder / "practice.db"

# 1. 读取 Excel：第一行是列名，后面的行才是数据。
workbook = load_workbook(excel_path, read_only=True, data_only=True)
try:
    sheet = workbook["tech_sites"]
    excel_rows = sheet.iter_rows(values_only=True)
    headers = next(excel_rows)
    expected_headers = ("id", "name", "category", "url", "description", "country")
    if headers != expected_headers:
        raise ValueError("Excel 表头与预期不同，请检查第一行。")
    rows = [row for row in excel_rows if any(value is not None for value in row)]
finally:
    workbook.close()

# 2. 连接数据库：文件不存在时，SQLite 会自动创建它。
connection = sqlite3.connect(db_path)
try:
    with connection:
        # 已导入过就停止，避免重复插入或覆盖你后来练习修改的数据。
        exists = connection.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
            ("tech_sites_practice",),
        ).fetchone()
        if exists:
            raise SystemExit("该表已经存在，无需再次导入。可以直接运行 SELECT 查询。")

        # 3. 创建表。INTEGER 表示整数，TEXT 表示文字。
        connection.execute("""
            CREATE TABLE tech_sites_practice (
                id INTEGER PRIMARY KEY,
                name TEXT,
                category TEXT,
                url TEXT,
                description TEXT,
                country TEXT
            )
        """)

        # 4. 批量插入。六个 ? 分别接收每行的六个值。
        connection.executemany(
            "INSERT INTO tech_sites_practice VALUES (?, ?, ?, ?, ?, ?)", rows
        )
    # with 块正常结束后，会自动提交（保存）数据。
finally:
    connection.close()

print(f"导入完成：{len(rows)} 条数据")
print(f"数据库：{db_path}")
print("表名：tech_sites_practice")
